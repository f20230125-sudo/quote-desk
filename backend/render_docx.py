"""The same quotation as a Word document, for customers whose procurement wants to mark it up."""

import io
from decimal import Decimal

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from backend import catalog
from backend.quote import Quote, pct

GREY = RGBColor(0x6B, 0x6B, 0x6B)


def money(x: Decimal) -> str:
    return f"{x:,.2f}"


def _p(doc, text="", size=10, bold=False, grey=False, align=None, space_after=2):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    if grey:
        run.font.color.rgb = GREY
    if align:
        para.alignment = align
    para.paragraph_format.space_after = Pt(space_after)
    return para


def render_docx(quote: Quote) -> bytes:
    q, cust, co, brand = quote.q, quote.customer, catalog.company(), catalog.assets()["brand"]
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.text = f"{co['bank']}    {q.quote_no} v{q.version}"
        footer.runs[0].font.size = Pt(7)

    head = _p(doc, "", space_after=0)
    r = head.add_run(brand["wordmark"])
    r.bold, r.font.size = True, Pt(20)
    r2 = head.add_run(f"  {brand['tagline']}")
    r2.font.size = Pt(11)
    _p(doc, f"{co['name']}, {co['address']}. TRN {co['trn']}", size=8, grey=True, space_after=10)
    if quote.blocked:
        _p(doc, "DRAFT: not to be sent until the checks pass.", size=10, bold=True, space_after=6)

    _p(doc, "Quotation", size=18, bold=True, space_after=6)
    meta = doc.add_table(rows=0, cols=2)
    for label, value in (
        ("Quotation", f"{q.quote_no} (version {q.version})"),
        ("Date", f"{q.issue_date:%d %B %Y}"),
        ("Valid until", f"{quote.valid_until:%d %B %Y}"),
        ("Your reference", q.their_ref or "-"),
        ("Prepared by", q.prepared_by or "-"),
        ("For", cust["name"]),
        ("Address", f"{cust['address']}, {cust['emirate']}"),
        ("TRN", cust.get("trn") or "not on file"),
    ):
        cells = meta.add_row().cells
        cells[0].text, cells[1].text = label, value
    _p(doc, "", space_after=6)

    table = doc.add_table(rows=1, cols=6)
    table.style = "Light Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, label in zip(
        table.rows[0].cells,
        ("#", "Part and description", "Qty", "Unit price", "Disc.", "Amount AED"),
        strict=True,
    ):
        cell.text = label
    for ln in quote.lines:
        cells = table.add_row().cells
        cells[0].text = str(ln.n)
        cells[1].text = f"{ln.part_number}\n{ln.description}"
        cells[2].text = f"{ln.qty} {ln.unit if ln.qty == 1 else ln.unit + 's'}"
        cells[3].text = money(ln.unit_price)
        cells[4].text = f"{pct(ln.discount_pct)}%" if ln.discount_pct else ""
        cells[5].text = money(ln.net)
        for k in (2, 3, 4, 5):
            cells[k].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for label, value in (("Subtotal", quote.subtotal), ("VAT 5%", quote.vat), ("Total AED", quote.total)):
        _p(doc, f"{label}   {money(value)}", bold=label.startswith("Total"), align=WD_ALIGN_PARAGRAPH.RIGHT)
    _p(
        doc,
        f"Longest lead time on this quotation: {quote.longest_lead} days from your purchase order.",
        size=8,
        grey=True,
        space_after=10,
    )
    if quote.clauses:
        _p(doc, "Terms", size=12, bold=True)
        for c in quote.clauses:
            para = _p(doc, "", space_after=4)
            t = para.add_run(f"{c['title']}. ")
            t.bold = True
            para.add_run(c["text"])
    _p(doc, "", space_after=16)
    _p(
        doc,
        f"For {co['name']}: ______________________        Accepted for {cust['name']}: ______________________",
    )
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
