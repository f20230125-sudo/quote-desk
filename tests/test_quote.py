"""Quotations: the figures, every check, and the PDF and Word files they become."""

import io
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader

from app import app
from backend.quote import QuoteIn, build, pct
from backend.render_docx import render_docx
from backend.render_pdf import render_pdf

TODAY = date(2026, 10, 22)
CLAUSES = ["payment", "validity", "delivery", "warranty", "einvoice", "law"]


def quote(**over) -> dict:
    base = dict(
        quote_no="Q-2026-0412",
        customer_code="GCB",
        issue_date="2026-10-22",
        prepared_by="Test",
        lines=[dict(part_number="PER760-XS2", qty=2, unit_price="45900.00")],
        clauses=CLAUSES,
    )
    base.update(over)
    return base


def messages(**over) -> list[tuple[str, str]]:
    q = build(QuoteIn.model_validate(quote(**over)), today=TODAY)
    return [(f.severity, f.message) for f in q.findings]


def test_a_clean_quote_has_nothing_to_say():
    assert messages() == []


def test_figures():
    q = build(
        QuoteIn.model_validate(
            quote(lines=[dict(part_number="PER760-XS2", qty=2, unit_price="45900.00", discount_pct="5")])
        ),
        today=TODAY,
    )
    ln = q.lines[0]
    assert ln.net == Decimal("87210.00")  # 45,900 less 5% = 43,605 each
    assert q.vat == Decimal("4360.50") and q.total == Decimal("91570.50")
    assert ln.margin == Decimal("87210.00") - Decimal("37635.00") * 2


def test_discount_is_rounded_after_multiplying_by_quantity():
    q = build(
        QuoteIn.model_validate(
            quote(lines=[dict(part_number="VVF-CORE-1Y", qty=16, unit_price="620.00", discount_pct="3.33")])
        ),
        today=TODAY,
    )
    assert q.lines[0].net == q.subtotal == Decimal("9589.66")
    assert q.vat == Decimal("479.48") and q.total == Decimal("10069.14")
    assert q.margin == Decimal("9589.66") - q.lines[0].cost * 16


def test_below_cost_blocks():
    msgs = messages(lines=[dict(part_number="SFP28-SR-25G", qty=4, unit_price="500.00")])
    assert (
        "block",
        "Line 1: SFP28-SR-25G is quoted below our cost (500.00 each against a cost of 552.00).",
    ) in msgs


def test_thin_margin_warns():
    msgs = messages(lines=[dict(part_number="SFP28-SR-25G", qty=4, unit_price="600.00")])
    assert ("warn", "Line 1: margin on SFP28-SR-25G is 8%, under the 10% floor.") in msgs


def test_margin_rounding_does_not_hide_a_policy_warning():
    q = build(
        QuoteIn.model_validate(quote(lines=[dict(part_number="SFP28-SR-25G", qty=1, unit_price="613.30")])),
        today=TODAY,
    )
    assert q.lines[0].margin_pct == Decimal("10.00")
    assert any(f.severity == "warn" and "floor" in f.message for f in q.findings)


def test_big_discount_needs_an_approver():
    over = dict(lines=[dict(part_number="PER760-XS2", qty=1, unit_price="48250.00", discount_pct="20")])
    assert any("needs the Sales Manager's approval" in m for _, m in messages(**over))
    assert not any("approval" in m for _, m in messages(**over, approved_by="R. Menon"))


def test_minimum_quantities():
    msgs = messages(lines=[dict(part_number="VVF-CORE-1Y", qty=8, unit_price="620.00")])
    assert ("block", "Line 1: VVF-CORE-1Y is sold in at least 16 licences; 8 cannot be ordered.") in msgs


def test_customer_without_trn():
    assert any("no TRN on file" in m for _, m in messages(customer_code="NWH"))


def test_expired_validity_blocks():
    assert ("block", "Validity ended on 21 Sep 2026. Reprice before sending.") in messages(
        issue_date="2026-08-22"
    )


def test_old_clause_version_warns_and_missing_clauses_too():
    msgs = [m for _, m in messages(clause_versions={"payment": 1}, clauses=["payment", "delivery"])]
    assert (
        'Clause "Payment" is version 1; version 2 is in force. Update it unless this quotation was agreed on the old terms.'
        in msgs
    )
    assert 'The "Validity" clause is not included.' in msgs


def test_clause_in_force_follows_the_date():
    q = build(QuoteIn.model_validate(quote(issue_date="2026-05-01")), today=date(2026, 5, 1))
    assert next(c for c in q.clauses if c["id"] == "payment")["v"] == 1


def test_unknown_customer_is_refused():
    with pytest.raises(ValueError):
        QuoteIn.model_validate(quote(customer_code="NOPE"))


def test_percent_wording():
    assert (pct(Decimal("20.00")), pct(Decimal("12.50")), pct(Decimal("0"))) == ("20", "12.5", "0")


def test_pdf_has_the_figures_and_marks_drafts():
    clean = build(QuoteIn.model_validate(quote()), today=TODAY)
    text = "".join(p.extract_text() for p in PdfReader(io.BytesIO(render_pdf(clean))).pages)
    assert "Q-2026-0412" in text and "96,390.00" in text and "Page 1 of 1" in text and "DRAFT" not in text
    blocked = build(QuoteIn.model_validate(quote(issue_date="2026-08-22")), today=TODAY)
    assert "DRAFT" in "".join(p.extract_text() for p in PdfReader(io.BytesIO(render_pdf(blocked))).pages)


def test_long_quotes_run_onto_more_pages():
    lines = [dict(part_number="P2425H-MON", qty=10 + i, unit_price="985.00") for i in range(40)]
    q = build(QuoteIn.model_validate(quote(lines=lines)), today=TODAY)
    reader = PdfReader(io.BytesIO(render_pdf(q)))
    assert len(reader.pages) > 1
    assert f"Page {len(reader.pages)} of {len(reader.pages)}" in reader.pages[-1].extract_text()


def test_word_file_has_the_lines_and_terms():
    q = build(QuoteIn.model_validate(quote()), today=TODAY)
    doc = Document(io.BytesIO(render_docx(q)))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells
    )
    assert "PER760-XS2" in text and "Payment within 30 days" in text and "Total AED" in text


@pytest.mark.parametrize("prepared_by", ["Uzair Khan", ""])
def test_word_file_includes_the_preparer(prepared_by):
    q = build(QuoteIn.model_validate(quote(prepared_by=prepared_by)), today=TODAY)
    doc = Document(io.BytesIO(render_docx(q)))
    rows = [(r.cells[0].text, r.cells[1].text) for r in doc.tables[0].rows]
    assert ("Prepared by", prepared_by or "-") in rows


def test_readme_runs_flask_from_the_installed_environment():
    readme = (Path(__file__).resolve().parents[1] / "README.md").read_text(encoding="utf-8")
    assert ".venv/Scripts/python -m flask --app app run -p 5201" in readme


@pytest.fixture
def client():
    app.config["TESTING"] = True
    return app.test_client()


@pytest.mark.parametrize("endpoint", ["check", "pdf", "docx"])
@pytest.mark.parametrize("version", [99, 0, -1])
def test_api_refuses_unknown_clause_versions(client, endpoint, version):
    r = client.post(f"/api/quote/{endpoint}", json=quote(clause_versions={"payment": version}))
    assert r.status_code == 400
    assert any(d["field"] == "clause_versions" for d in r.get_json()["error"]["details"])


@pytest.mark.parametrize("endpoint", ["check", "pdf", "docx"])
@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("quote_no", "Q-TEST\r\nX-Demo: yes"),
        ("quote_no", "Q-TEST\n"),
        ("quote_no", "Q-TEST\r"),
        ("quote_no", "Q-TEST\x00"),
        ("their_ref", "RFQ\x00X"),
        ("their_ref", "RFQ\x0bX"),
        ("their_ref", "RFQ\ufffeX"),
        ("prepared_by", "Name\x0cX"),
        ("approved_by", "Name\x1fX"),
    ],
)
def test_api_refuses_characters_that_break_documents(client, endpoint, field, value):
    r = client.post(f"/api/quote/{endpoint}", json=quote(**{field: value}))
    assert r.status_code == 400
    assert any(d["field"] == field for d in r.get_json()["error"]["details"])


@pytest.mark.parametrize("endpoint", ["check", "pdf", "docx"])
@pytest.mark.parametrize("validity", [{}, {"valid_days": 30}])
def test_api_refuses_an_expiry_beyond_the_calendar(client, endpoint, validity):
    r = client.post(f"/api/quote/{endpoint}", json=quote(issue_date="9999-12-31", **validity))
    assert r.status_code == 400
    assert any(d["field"] == "valid_days" for d in r.get_json()["error"]["details"])


def test_api_accepts_the_last_representable_expiry(client):
    r = client.post("/api/quote/check", json=quote(issue_date="9999-12-30", valid_days=1))
    assert r.status_code == 200 and r.get_json()["valid_until"] == "9999-12-31"


def test_api_check_and_files(client):
    r = client.post("/api/quote/check", json=quote())
    assert r.status_code == 200 and r.get_json()["total"] == "96390.00"
    assert client.post("/api/quote/pdf", json=quote()).data[:5] == b"%PDF-"
    assert client.post("/api/quote/docx", json=quote()).data[:2] == b"PK"
    bad = client.post("/api/quote/check", json=quote(lines=[dict(part_number="X", qty=0, unit_price="-1")]))
    assert bad.status_code == 400 and {d["field"] for d in bad.get_json()["error"]["details"]} >= {
        "lines.0.qty",
        "lines.0.unit_price",
    }
    assert client.get("/api/data").get_json()["assets"]["policy"]["margin_floor_pct"] == "10"
    assert client.get("/api/nope").status_code == 404
